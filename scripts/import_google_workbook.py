"""Import every game tab from the Google Sheets stat workbook into volleyball_stats.xlsx."""
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = ROOT / "source" / "varsity_stat_workbook.xlsx"
XLSX = ROOT / "volleyball_stats.xlsx"
DOWNLOADS = Path.home() / "Downloads" / "volleyball_stats.xlsx"
SKIP_TABS = {"copy - do not fill"}


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_tab_name(name):
    if text(name).lower() in SKIP_TABS or text(name).lower().startswith("copy"):
        return None
    base = text(name)
    match = re.search(r"(\d{2,4})\s*$", base.replace("-", " "))
    if not match:
        return None
    digits = match.group(1)
    month_day = split_month_day(digits)
    if not month_day:
        return None
    month, day = month_day
    opponent = re.sub(r"[\s\-_]+(\d{2,4})\s*$", "", base).strip(" -_")
    opponent = re.sub(r"\s+", " ", opponent)
    if opponent.lower() == "heritage chrisian":
        opponent = "Heritage Christian"
    date = datetime(2026, month, day)
    return date, opponent


def split_month_day(digits):
    """Month and day are concatenated with no separator and no leading zeros.

    821 is August 21, 91 is September 1, and 107 is October 7. A 3-digit
    value starting with 10, 11, or 12 is that month plus a one-digit day.
    Reading 107 as January 7 only works if the day is written with a leading
    zero, which these tab names never do.
    """
    if len(digits) == 4:
        month, day = int(digits[:2]), int(digits[2:])
    elif len(digits) == 3 and digits.startswith(("10", "11", "12")):
        month, day = int(digits[:2]), int(digits[2])
    elif len(digits) == 3:
        month, day = int(digits[0]), int(digits[1:])
    elif len(digits) == 2:
        month, day = int(digits[0]), int(digits[1])
    else:
        return None
    if not (1 <= month <= 12 and 1 <= day <= 31):
        return None
    return month, day


def parse_digits(value):
    if value is None or value == "":
        return []
    return [int(part) for part in re.findall(r"\d+", str(value))]


def scores_string(nums):
    if not nums:
        return None
    return " ".join(str(n) for n in nums)


def num(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    token = str(value).replace("%", "").strip()
    if not token:
        return None
    if "." in token:
        return float(token)
    return int(token)


def derived_rating(rating, attempts, parts):
    """Keep a rating typed on the sheet. If that cell is blank, divide the recorded grades by the attempts."""
    if rating is not None:
        return rating
    if not parts or not attempts:
        return None
    return round(sum(parts) / float(attempts), 2)


def header_map(ws, row_index):
    mapping = {}
    for col in range(1, ws.max_column + 1):
        label = text(ws.cell(row_index, col).value)
        if label:
            mapping[col] = label
    return mapping


def col_by_label(header, contains):
    contains = [part.lower() for part in contains]
    for col, label in header.items():
        lower = label.lower()
        if all(part in lower for part in contains):
            return col
    return None


def row_value(ws, row, col):
    if col is None:
        return None
    return ws.cell(row, col).value


def find_row(ws, start, predicate):
    for row in range(start, ws.max_row + 1):
        if predicate(row):
            return row
    return None


def serving_team_row(ws):
    return find_row(
        ws,
        2,
        lambda row: text(ws.cell(row, 1).value).lower() == "team"
        and "total aces" in text(ws.cell(row, 2).value).lower(),
    )


def is_player_name(name):
    if not name:
        return False
    lower = name.lower().strip()
    if lower in {"player", "team", "total"}:
        return False
    if lower.startswith("rotation") or lower.startswith("set "):
        return False
    if lower.startswith("team kill") or lower.startswith("team pass"):
        return False
    if re.match(r"^-?\d+(\.\d+)?$", name.strip()):
        return False
    return bool(re.match(r"^[A-Za-z]", name.strip()))


def player_section_end(ws):
    stops = []
    rotation_row = find_row(
        ws,
        2,
        lambda row: "rotation" in text(ws.cell(row, 1).value).lower()
        and "setter" in text(ws.cell(row, 1).value).lower(),
    )
    if rotation_row:
        stops.append(rotation_row)
    team_row = serving_team_row(ws)
    if team_row:
        stops.append(team_row)
    if not stops:
        return ws.max_row + 1
    return min(stops)


def player_rows(ws, date, opponent):
    section_end = player_section_end(ws)
    header = header_map(ws, 1)
    cols = {
        "digs": col_by_label(header, ["dig"]),
        "assists": col_by_label(header, ["assist"]),
        "attacks": col_by_label(header, ["attack attempt"]),
        "kills": None,
        "kill_errors": col_by_label(header, ["kill error"]),
        "kill_eff": col_by_label(header, ["kill efficiency"]),
        "serve_attempts": col_by_label(header, ["serve attempt"]),
        "aces": col_by_label(header, ["ace"]),
        "serve_errors": col_by_label(header, ["serving error"]),
        "serve_scores": col_by_label(header, ["serving scores"]),
        "serve_rating": col_by_label(header, ["serving rating"]),
        "unforced": col_by_label(header, ["unforced error"]),
        "stuff": col_by_label(header, ["block (stuff)"]),
        "touch": col_by_label(header, ["block (touch)"]),
        "recv_attempts": col_by_label(header, ["serve receive attempt"]),
        "recv_grades": col_by_label(header, ["serve receive grade"]),
        "recv_rating": col_by_label(header, ["serve receive rating"]),
        "fb_attempts": col_by_label(header, ["fb pass attempt"]),
        "fb_grades": col_by_label(header, ["fb pass grade"]),
        "fb_rating": col_by_label(header, ["fb receive rating"]),
    }
    for col, label in header.items():
        lower = label.lower()
        if lower == "kill":
            cols["kills"] = col
        elif lower.startswith("kill error"):
            cols["kill_errors"] = col
        elif "kill efficiency" in lower:
            cols["kill_eff"] = col

    rows = []
    for row in range(2, section_end):
        name = text(ws.cell(row, 1).value)
        if not is_player_name(name):
            continue
        # Entered on the Oct 7 First Baptist tab, but she did not play that match.
        if (
            name == "Bridget"
            and opponent == "First Baptist"
            and date.month == 10
            and date.day == 7
        ):
            continue
        serve_scores = parse_digits(row_value(ws, row, cols["serve_scores"]))
        recv_grades = parse_digits(row_value(ws, row, cols["recv_grades"]))
        fb_grades = parse_digits(row_value(ws, row, cols["fb_grades"]))
        serve_attempts = num(row_value(ws, row, cols["serve_attempts"]))
        receive_attempts = num(row_value(ws, row, cols["recv_attempts"]))
        freeball_attempts = num(row_value(ws, row, cols["fb_attempts"]))
        rows.append(
            {
                "date": date,
                "opponent": opponent,
                "player": name,
                "digs": num(row_value(ws, row, cols["digs"])),
                "assists": num(row_value(ws, row, cols["assists"])),
                "attack_attempts": num(row_value(ws, row, cols["attacks"])),
                "kills": num(row_value(ws, row, cols["kills"])),
                "kill_errors": num(row_value(ws, row, cols["kill_errors"])),
                "kill_efficiency": num(row_value(ws, row, cols["kill_eff"])),
                "serve_attempts": serve_attempts,
                "aces": num(row_value(ws, row, cols["aces"])),
                "serve_errors": num(row_value(ws, row, cols["serve_errors"])),
                "serve_rating": derived_rating(
                    num(row_value(ws, row, cols["serve_rating"])),
                    serve_attempts,
                    serve_scores,
                ),
                "serving_scores": scores_string(serve_scores),
                "unforced_errors": num(row_value(ws, row, cols["unforced"])),
                "stuff_blocks": num(row_value(ws, row, cols["stuff"])),
                "block_touches": num(row_value(ws, row, cols["touch"])),
                "serve_receive_attempts": receive_attempts,
                "serve_receive_rating": derived_rating(
                    num(row_value(ws, row, cols["recv_rating"])),
                    receive_attempts,
                    recv_grades,
                ),
                "serve_receive_grades": scores_string(recv_grades),
                "freeball_attempts": freeball_attempts,
                "freeball_rating": derived_rating(
                    num(row_value(ws, row, cols["fb_rating"])),
                    freeball_attempts,
                    fb_grades,
                ),
                "freeball_grades": scores_string(fb_grades),
            }
        )
    return rows


def rotation_rows(ws, date, opponent):
    header_row = find_row(
        ws,
        2,
        lambda row: "rotation" in text(ws.cell(row, 1).value).lower()
        and "setter" in text(ws.cell(row, 1).value).lower(),
    )
    if header_row is None:
        return []
    team_row = find_row(
        ws,
        header_row + 1,
        lambda row: text(ws.cell(row, 1).value).lower() == "team"
        and "total aces" in text(ws.cell(row, 2).value).lower(),
    )
    if team_row is None:
        team_row = ws.max_row + 1
    header = header_map(ws, header_row)
    cols = {
        "assists": col_by_label(header, ["assist"]),
        "attacks": col_by_label(header, ["attack attempt"]),
        "kills": None,
        "kill_errors": col_by_label(header, ["kill error"]),
        "kill_eff": col_by_label(header, ["kill efficiency"]),
        "unforced": col_by_label(header, ["unforced error"]),
        "recv_attempts": col_by_label(header, ["serve receive attempt"]),
        "recv_avg": col_by_label(header, ["serve receive avg"]),
        "serve_aces": col_by_label(header, ["serve ace"]),
        "serve_errors": col_by_label(header, ["serve error"]),
        "points_for": col_by_label(header, ["points scored in rotation"]),
        "points_against": col_by_label(header, ["points scored against"]),
        "net": col_by_label(header, ["rotation net"]),
    }
    for col, label in header.items():
        if label.lower() == "kill":
            cols["kills"] = col

    rows = []
    for row in range(header_row + 1, team_row):
        rot = ws.cell(row, 1).value
        if rot is None:
            continue
        if isinstance(rot, str) and not rot.strip():
            continue
        try:
            rotation = int(float(rot))
        except (TypeError, ValueError):
            continue
        if rotation < 1 or rotation > 6:
            continue
        rows.append(
            {
                "date": date,
                "opponent": opponent,
                "rotation": rotation,
                "assists": num(row_value(ws, row, cols["assists"])),
                "attack_attempts": num(row_value(ws, row, cols["attacks"])),
                "kills": num(row_value(ws, row, cols["kills"])),
                "kill_errors": num(row_value(ws, row, cols["kill_errors"])),
                "kill_efficiency": num(row_value(ws, row, cols["kill_eff"])),
                "unforced_errors": num(row_value(ws, row, cols["unforced"])),
                "serve_receive_attempts": num(row_value(ws, row, cols["recv_attempts"])),
                "serve_receive_avg": num(row_value(ws, row, cols["recv_avg"])),
                "serve_aces": num(row_value(ws, row, cols["serve_aces"])),
                "serve_errors": num(row_value(ws, row, cols["serve_errors"])),
                "points_for": num(row_value(ws, row, cols["points_for"])),
                "points_against": num(row_value(ws, row, cols["points_against"])),
                "net": num(row_value(ws, row, cols["net"])),
            }
        )
    return rows


def serving_rows(ws, date, opponent):
    team_row = serving_team_row(ws)
    if team_row is None:
        return []
    rows = []
    for row in range(team_row + 1, ws.max_row + 1):
        label = text(ws.cell(row, 1).value)
        if not label:
            continue
        lower = label.lower()
        if lower.startswith("team kill efficiency"):
            break
        if lower.startswith("set "):
            set_no = int(re.search(r"\d+", label).group())
            rows.append(
                {
                    "date": date,
                    "opponent": opponent,
                    "set": set_no,
                    "aces_us": num(ws.cell(row, 2).value),
                    "errors_us": num(ws.cell(row, 3).value),
                    "ace_pct_us": num(ws.cell(row, 4).value),
                    "error_pct_us": num(ws.cell(row, 5).value),
                    "aces_them": num(ws.cell(row, 6).value),
                    "errors_them": num(ws.cell(row, 7).value),
                    "ace_pct_them": num(ws.cell(row, 8).value),
                    "error_pct_them": num(ws.cell(row, 9).value),
                }
            )
        elif lower == "total":
            rows.append(
                {
                    "date": date,
                    "opponent": opponent,
                    "set": "Total",
                    "aces_us": num(ws.cell(row, 2).value),
                    "errors_us": num(ws.cell(row, 3).value),
                    "ace_pct_us": num(ws.cell(row, 4).value),
                    "error_pct_us": num(ws.cell(row, 5).value),
                    "aces_them": num(ws.cell(row, 6).value),
                    "errors_them": num(ws.cell(row, 7).value),
                    "ace_pct_them": num(ws.cell(row, 8).value),
                    "error_pct_them": num(ws.cell(row, 9).value),
                }
            )
    return rows


def match_row(ws, date, opponent):
    team_hitting = None
    note = ""
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            label = text(ws.cell(row, col).value).lower()
            if "team kill efficiency" not in label:
                continue
            for candidate in (
                ws.cell(row, col + 1).value,
                ws.cell(row + 1, col).value,
                ws.cell(row + 1, col + 1).value,
                ws.cell(row, 2).value,
            ):
                value = num(candidate)
                if value is not None and abs(value) <= 1.5:
                    team_hitting = value
                    break
            break
        if team_hitting is not None:
            break
    attack_row = find_row(
        ws,
        2,
        lambda row: "attack attempt" in text(ws.cell(row, 1).value).lower()
        and "all games" in text(ws.cell(row, 1).value).lower(),
    )
    if attack_row:
        note = text(ws.cell(attack_row, 2).value)
    return {
        "date": date,
        "opponent": opponent,
        "team_kill_efficiency": team_hitting,
        "note": note,
    }


def calc_kill_eff(row):
    attempts = row.get("attack_attempts")
    if not attempts:
        return None
    kills = row.get("kills") or 0
    errors = row.get("kill_errors") or 0
    return round((kills - errors) / float(attempts), 3)


def parse_game_tab(ws, tab_name):
    parsed = parse_tab_name(tab_name)
    if not parsed:
        return None
    date, opponent = parsed
    return {
        "tab": tab_name,
        "date": date,
        "opponent": opponent,
        "players": player_rows(ws, date, opponent),
        "rotations": rotation_rows(ws, date, opponent),
        "serving": serving_rows(ws, date, opponent),
        "match": match_row(ws, date, opponent),
    }


def write_sheet(ws, headers, rows):
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    for row in rows:
        ws.append(row)


def import_workbook(src_path):
    src = Path(src_path)
    wb_src = load_workbook(src, data_only=True)
    games = []
    for name in wb_src.sheetnames:
        game = parse_game_tab(wb_src[name], name)
        if game:
            games.append(game)
    games.sort(key=lambda item: (item["date"], item["opponent"], item["tab"]))

    players_out = []
    rotations_out = []
    serving_out = []
    matches_out = []
    for game in games:
        date, opponent = game["date"], game["opponent"]
        for row in game["players"]:
            players_out.append(
                [
                    date,
                    opponent,
                    row["player"],
                    row["digs"],
                    row["assists"],
                    row["attack_attempts"],
                    row["kills"],
                    row["kill_errors"],
                    row["kill_efficiency"],
                    calc_kill_eff(row),
                    row["serve_attempts"],
                    row["aces"],
                    row["serve_errors"],
                    row["serve_rating"],
                    row["serving_scores"],
                    row["unforced_errors"],
                    row["stuff_blocks"],
                    row["block_touches"],
                    row["serve_receive_attempts"],
                    row["serve_receive_rating"],
                    row["serve_receive_grades"],
                    row["freeball_attempts"],
                    row["freeball_rating"],
                    row["freeball_grades"],
                ]
            )
        for row in game["rotations"]:
            rotations_out.append(
                [
                    date,
                    opponent,
                    row["rotation"],
                    row["assists"],
                    row["attack_attempts"],
                    row["kills"],
                    row["kill_errors"],
                    row["kill_efficiency"],
                    row["unforced_errors"],
                    row["serve_receive_attempts"],
                    row["serve_receive_avg"],
                    row["serve_aces"],
                    row["serve_errors"],
                    row["points_for"],
                    row["points_against"],
                    row["net"],
                ]
            )
        for row in game["serving"]:
            serving_out.append(
                [
                    date,
                    opponent,
                    row["set"],
                    row["aces_us"],
                    row["errors_us"],
                    row["ace_pct_us"],
                    row["error_pct_us"],
                    row["aces_them"],
                    row["errors_them"],
                    row["ace_pct_them"],
                    row["error_pct_them"],
                ]
            )
        match = game["match"]
        matches_out.append([date, opponent, match["team_kill_efficiency"], match["note"] or None])

    wb = load_workbook(XLSX)
    write_sheet(
        wb["Player stats"],
        [
            "Date",
            "Opponent",
            "Player",
            "Digs",
            "Assists",
            "Attack attempts",
            "Kills",
            "Kill errors",
            "Kill efficiency",
            "Kill efficiency (calculated)",
            "Serve attempts",
            "Aces",
            "Serve errors",
            "Serve rating",
            "Serving scores",
            "Unforced errors",
            "Stuff blocks",
            "Block touches",
            "Serve receive attempts",
            "Serve receive rating",
            "Serve receive grades",
            "Freeball attempts",
            "Freeball rating",
            "Freeball grades",
        ],
        players_out,
    )
    write_sheet(
        wb["Rotation stats"],
        [
            "Date",
            "Opponent",
            "Rotation",
            "Assists",
            "Attack attempts",
            "Kills",
            "Kill errors",
            "Kill efficiency",
            "Unforced errors",
            "Serve receive attempts",
            "Serve receive average",
            "Serve aces",
            "Serve errors",
            "Points scored",
            "Points against",
            "Rotation net",
        ],
        rotations_out,
    )
    write_sheet(
        wb["Serving by set"],
        [
            "Date",
            "Opponent",
            "Set",
            "Aces (us)",
            "Serve errors (us)",
            "Ace % (us)",
            "Serve error % (us)",
            "Aces (them)",
            "Serve errors (them)",
            "Ace % (them)",
            "Serve error % (them)",
        ],
        serving_out,
    )
    write_sheet(
        wb["Matches"],
        ["Date", "Opponent", "Team kill efficiency", "Note"],
        matches_out,
    )
    wb.save(XLSX)
    if DOWNLOADS.parent.is_dir():
        shutil.copyfile(XLSX, DOWNLOADS)
    print("Imported {} game tabs from {}".format(len(games), src.name))
    for game in games:
        print(
            " ",
            game["date"].date().isoformat(),
            game["opponent"],
            "players",
            len(game["players"]),
            "rotations",
            len(game["rotations"]),
            "sets",
            len(game["serving"]),
        )


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    if not src.exists():
        raise SystemExit("Workbook not found: " + str(src))
    import_workbook(src)


if __name__ == "__main__":
    main()
