"""Turn volleyball_stats.xlsx into data.js for the stats explorer."""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "volleyball_stats.xlsx"
OUT = ROOT / "data.js"


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value


def scores(value):
    if not value:
        return []
    return [int(part) for part in str(value).split()]


def sheet_rows(ws):
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        rows.append({header: iso(value) for header, value in zip(headers, row)})
    return rows


def main():
    wb = load_workbook(XLSX, data_only=True)
    players = []
    for row in sheet_rows(wb["Player stats"]):
        players.append(
            {
                "date": row["Date"],
                "opponent": row["Opponent"],
                "player": row["Player"],
                "digs": row["Digs"],
                "assists": row["Assists"],
                "attackAttempts": row["Attack attempts"],
                "kills": row["Kills"],
                "killErrors": row["Kill errors"],
                "hittingSheet": row["Kill efficiency"],
                "serveAttempts": row["Serve attempts"],
                "aces": row["Aces"],
                "serveErrors": row["Serve errors"],
                "serveRating": row["Serve rating"],
                "serveScores": scores(row["Serving scores"]),
                "unforcedErrors": row["Unforced errors"],
                "stuffBlocks": row["Stuff blocks"],
                "blockTouches": row["Block touches"],
                "receiveAttempts": row["Serve receive attempts"],
                "receiveRating": row["Serve receive rating"],
                "receiveGrades": scores(row["Serve receive grades"]),
                "freeballAttempts": row["Freeball attempts"],
                "freeballRating": row["Freeball rating"],
                "freeballGrades": scores(row["Freeball grades"]),
            }
        )

    rotations = []
    for row in sheet_rows(wb["Rotation stats"]):
        rotations.append(
            {
                "date": row["Date"],
                "opponent": row["Opponent"],
                "rotation": row["Rotation"],
                "assists": row["Assists"],
                "attackAttempts": row["Attack attempts"],
                "kills": row["Kills"],
                "killErrors": row["Kill errors"],
                "unforcedErrors": row["Unforced errors"],
                "receiveAttempts": row["Serve receive attempts"],
                "receiveAvg": row["Serve receive average"],
                "serveAces": row["Serve aces"],
                "serveErrors": row["Serve errors"],
                "pointsFor": row["Points scored"],
                "pointsAgainst": row["Points against"],
                "net": row["Rotation net"],
            }
        )

    serving = []
    for row in sheet_rows(wb["Serving by set"]):
        serving.append(
            {
                "date": row["Date"],
                "opponent": row["Opponent"],
                "set": row["Set"],
                "acesUs": row["Aces (us)"],
                "errorsUs": row["Serve errors (us)"],
                "acePctUs": row["Ace % (us)"],
                "errorPctUs": row["Serve error % (us)"],
                "acesThem": row["Aces (them)"],
                "errorsThem": row["Serve errors (them)"],
                "acePctThem": row["Ace % (them)"],
                "errorPctThem": row["Serve error % (them)"],
            }
        )

    matches = []
    for row in sheet_rows(wb["Matches"]):
        matches.append(
            {
                "date": row["Date"],
                "opponent": row["Opponent"],
                "teamHitting": row["Team kill efficiency"],
                "note": row["Note"] or "",
            }
        )

    payload = {
        "players": players,
        "rotations": rotations,
        "serving": serving,
        "matches": matches,
    }
    OUT.write_text(
        "window.DATA = " + json.dumps(payload, indent=2) + ";\n",
        encoding="utf-8",
    )
    print("wrote {} ({} players, {} matches)".format(OUT, len(players), len(matches)))


if __name__ == "__main__":
    main()
