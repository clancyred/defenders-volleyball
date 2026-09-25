"""Parse stat-sheet PDFs and add any new matches to volleyball_stats.xlsx."""
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_stat_sheet import calculated_efficiency, checked_sequence, parse_pdf

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "volleyball_stats.xlsx"
DOWNLOADS = Path.home() / "Downloads" / "volleyball_stats.xlsx"
TOTAL_FILL = PatternFill("solid", fgColor="D6E3F0")
BOLD = Font(bold=True)


def as_date(value):
    if isinstance(value, datetime):
        return value
    year, month, day = [int(part) for part in value.split("-")]
    return datetime(year, month, day)


def pct(value):
    if value is None:
        return None
    return value / 100.0


def player_row(row):
    return [
        as_date(row["date"]),
        row["opponent"],
        row["player"],
        row.get("digs"),
        row.get("assists"),
        row.get("attack_attempts"),
        row.get("kills"),
        row.get("kill_errors"),
        row.get("kill_efficiency"),
        calculated_efficiency(row),
        row.get("serve_attempts"),
        row.get("aces"),
        row.get("serve_errors"),
        row.get("serve_rating"),
        checked_sequence(row.get("serving_scores"), row.get("serve_attempts"), row.get("serve_rating"), 0, 5),
        row.get("unforced_errors"),
        row.get("blocks_stuff"),
        row.get("blocks_touch"),
        row.get("serve_receive_attempts"),
        row.get("serve_receive_rating"),
        checked_sequence(row.get("sr_grades"), row.get("serve_receive_attempts"), row.get("serve_receive_rating"), 0, 3),
        row.get("fb_pass_attempts"),
        row.get("fb_receive_rating"),
        checked_sequence(row.get("fb_grades"), row.get("fb_pass_attempts"), row.get("fb_receive_rating"), 0, 3),
    ]


def rotation_row(row):
    return [
        as_date(row["date"]),
        row["opponent"],
        row.get("rotation"),
        row.get("assists"),
        row.get("attack_attempts"),
        row.get("kills"),
        row.get("kill_errors"),
        row.get("kill_efficiency"),
        row.get("unforced_errors"),
        row.get("serve_receive_attempts"),
        row.get("serve_receive_avg"),
        row.get("serve_aces"),
        row.get("serve_errors"),
        row.get("points_scored"),
        row.get("points_against"),
        row.get("net_score"),
    ]


def serving_row(row):
    return [
        as_date(row["date"]),
        row["opponent"],
        row.get("set"),
        row.get("aces_us"),
        row.get("serve_errors_us"),
        pct(row.get("ace_pct_us")),
        pct(row.get("serve_error_pct_us")),
        row.get("aces_them"),
        row.get("serve_errors_them"),
        pct(row.get("ace_pct_them")),
        pct(row.get("serve_error_pct_them")),
    ]


def match_row(match):
    return [
        as_date(match["date"]),
        match["opponent"],
        match.get("team_kill_efficiency"),
        match.get("note") or None,
    ]


def existing_rows(ws):
    rows = []
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if any(value is not None for value in row):
            rows.append((list(row), index))
    return rows


def write_rows(ws, rows, formats, total_column=None):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for values, _seq in rows:
        ws.append(values)
        excel_row = ws.max_row
        for col, fmt in formats.items():
            if ws.cell(excel_row, col).value is not None:
                ws.cell(excel_row, col).number_format = fmt
        if total_column and values[total_column - 1] == "Total":
            for col in range(1, len(values) + 1):
                cell = ws.cell(excel_row, col)
                cell.font = BOLD
                cell.fill = TOTAL_FILL


def main(paths):
    wb = load_workbook(XLSX)
    known = set()
    for row in wb["Matches"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        known.add((as_date(row[0]).date().isoformat(), row[1]))

    new_players, new_rots, new_sets, new_matches = [], [], [], []
    for path in paths:
        parsed = parse_pdf(path)
        match = parsed["match"]
        key = (match["date"], match["opponent"])
        if key in known:
            print("already in workbook:", key)
            continue
        print("adding", match["date"], match["opponent"], "players", len(parsed["players"]), "issues", len(parsed["issues"]))
        for issue in parsed["issues"]:
            print(" ", issue)
        new_players.extend(player_row(row) for row in parsed["players"])
        new_rots.extend(rotation_row(row) for row in parsed["rotations"])
        new_sets.extend(serving_row(row) for row in parsed["sets"])
        new_sets.append(serving_row(parsed["total"]))
        new_matches.append(match_row(match))
        known.add(key)

    def combine(ws, extra, sort_key):
        rows = existing_rows(ws)
        start = len(rows)
        rows.extend((values, start + index) for index, values in enumerate(extra))
        rows.sort(key=sort_key)
        return rows

    def date_opp(item):
        values, seq = item
        when = as_date(values[0])
        return (when, values[1], seq)

    def date_opp_set(item):
        values, seq = item
        when = as_date(values[0])
        set_no = 99 if values[2] == "Total" else int(values[2])
        return (when, values[1], set_no, seq)

    write_rows(
        wb["Player stats"],
        combine(wb["Player stats"], new_players, date_opp),
        {1: "YYYY-MM-DD", 9: "0.000", 10: "0.000", 14: "0.00", 20: "0.00", 23: "0.00"},
    )
    write_rows(
        wb["Rotation stats"],
        combine(wb["Rotation stats"], new_rots, date_opp),
        {1: "YYYY-MM-DD", 8: "0.000", 11: "0.00"},
    )
    write_rows(
        wb["Serving by set"],
        combine(wb["Serving by set"], new_sets, date_opp_set),
        {1: "YYYY-MM-DD", 6: "0.00%", 7: "0.00%", 10: "0.00%", 11: "0.00%"},
        total_column=3,
    )
    write_rows(
        wb["Matches"],
        combine(wb["Matches"], new_matches, date_opp),
        {1: "YYYY-MM-DD", 3: "0.000"},
    )
    wb.save(XLSX)
    shutil.copyfile(XLSX, DOWNLOADS)
    print("saved", XLSX)


if __name__ == "__main__":
    main(sys.argv[1:])
